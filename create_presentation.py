from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
OUTPUT = ROOT / 'energy_demand_forecast_presentation.pptx'
FONT = 'Apple SD Gothic Neo'
NAVY = RGBColor(15, 31, 48)
BLUE = RGBColor(32, 91, 137)
CYAN = RGBColor(67, 190, 202)
ORANGE = RGBColor(240, 151, 70)
PALE = RGBColor(239, 245, 248)
MID = RGBColor(96, 117, 132)
DARK = RGBColor(31, 43, 53)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(75, 157, 111)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
metrics = pd.read_csv(ROOT / 'forecast_metrics.csv')


def set_bg(slide, color=WHITE):
    shape = slide.background
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def text(slide, value, x, y, w, h, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(2)
    frame.margin_right = Pt(2)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def title(slide, heading, kicker=None, heading_color=NAVY):
    if kicker:
        text(slide, kicker.upper(), 0.65, 0.28, 5.8, 0.25, 8, CYAN, True)
    text(slide, heading, 0.65, 0.58, 12, 0.55, 25, heading_color, True)
    rect(slide, 0.68, 1.25, 1.0, 0.05, CYAN)


def footer(slide, number):
    text(slide, 'ENERGY DEMAND FORECASTING  /  2025', 0.65, 7.16, 5, 0.18, 7, MID, True)
    text(slide, str(number).zfill(2), 12.1, 7.13, 0.55, 0.22, 8, MID, True, PP_ALIGN.RIGHT)


def bullet_list(slide, items, x, y, w, size=16, color=DARK, gap=0.48):
    for index, item in enumerate(items):
        yy = y + index * gap
        rect(slide, x, yy + 0.13, 0.09, 0.09, CYAN)
        text(slide, item, x + 0.22, yy, w - 0.22, 0.36, size, color)


def add_image(slide, filename, x, y, w, h=None):
    path = ROOT / filename
    if h is None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def metric_card(slide, x, y, label, value, accent):
    rect(slide, x, y, 2.85, 1.15, PALE, PALE, True)
    rect(slide, x, y, 0.08, 1.15, accent)
    text(slide, label, x + 0.22, y + 0.18, 2.35, 0.25, 10, MID, True)
    text(slide, value, x + 0.22, y + 0.48, 2.35, 0.42, 22, NAVY, True)


# 1. Cover
slide = prs.slides.add_slide(blank)
set_bg(slide, NAVY)
rect(slide, 8.85, 0, 4.5, 7.5, BLUE)
rect(slide, 9.55, 0.85, 2.9, 0.08, CYAN)
text(slide, 'ENERGY / TIME SERIES / 2025', 0.8, 0.85, 5, 0.3, 11, CYAN, True)
text(slide, '전국 전력수요\n시계열 예측', 0.8, 1.55, 7.5, 1.45, 34, WHITE, True)
text(slide, 'ARIMA · LSTM · Transformer 비교\n7개 지역 ASOS 기상 데이터 기반', 0.85, 3.35, 6.6, 0.7, 17, RGBColor(215, 230, 237))
text(slide, '주희', 0.85, 5.95, 2, 0.35, 14, WHITE, True)
text(slide, 'National Power Demand and Regional Weather in 2025', 0.85, 6.42, 7.2, 0.3, 10, RGBColor(180, 207, 219))
for i, (yy, ww, cc) in enumerate([(1.4, 2.0, CYAN), (2.3, 2.8, ORANGE), (3.2, 1.65, WHITE), (4.1, 2.5, CYAN), (5.0, 1.85, ORANGE)]):
    rect(slide, 9.4, yy, ww, 0.16, cc, cc, True)

# 2. Question
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, '왜 전력수요를 예측하는가?', '01 / problem')
text(slide, '전력수요는 시간대와 날씨에 따라 반복적으로 변하지만,\n그 변동의 크기와 방향은 매일 다릅니다.', 0.8, 1.7, 7.2, 1.0, 24, NAVY, True)
bullet_list(slide, ['전국 수요의 다음 1시간 값을 예측', '지역별 기상 변수를 하나의 다변량 입력으로 사용', '고전 시계열 모델과 딥러닝 모델의 성능 비교', '예측값뿐 아니라 불확실성과 변동성도 검증'], 0.95, 3.25, 6.7, 17, DARK, 0.58)
rect(slide, 8.25, 1.65, 4.15, 4.5, PALE, PALE, True)
text(slide, '연구 질문', 8.7, 2.1, 3.2, 0.3, 12, BLUE, True)
text(slide, '어떤 모델이\n실제 수요의 변동성을\n가장 잘 따라가는가?', 8.7, 2.75, 3.2, 1.55, 25, NAVY, True)
text(slide, '정확도 + 공정한 시험 + uncertainty', 8.7, 5.2, 3.0, 0.45, 12, ORANGE, True)
footer(slide, 2)

# 3. Data
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, '데이터: 전국 수요 + 7개 지역 날씨', '02 / data')
regions = ['서울', '인천', '대전', '대구', '광주', '부산', '제주']
for i, region in enumerate(regions):
    x = 0.85 + (i % 4) * 2.95
    y = 1.75 + (i // 4) * 1.0
    rect(slide, x, y, 2.45, 0.62, PALE, PALE, True)
    text(slide, region, x, y + 0.1, 2.45, 0.3, 16, NAVY, True, PP_ALIGN.CENTER)
text(slide, '통합 데이터셋', 0.85, 4.1, 2.4, 0.3, 12, BLUE, True)
metric_card(slide, 0.85, 4.55, '관측 시간', '8,760시간', CYAN)
metric_card(slide, 3.95, 4.55, '입력 특성', '50개', ORANGE)
metric_card(slide, 7.05, 4.55, '예측 목표', '전국 MWh', GREEN)
text(slide, 'national_power_regional_weather_2025.csv', 0.9, 6.25, 6.5, 0.35, 14, NAVY, True)
text(slide, '전처리된 전력수요·지역 날씨를 하나의 CSV로 저장해 재현성을 높였습니다.', 0.9, 6.62, 8.5, 0.25, 11, MID)
footer(slide, 3)

# 4. Preprocessing
slide = prs.slides.add_slide(blank)
set_bg(slide, WHITE)
title(slide, '전처리와 데이터 누출 방지', '03 / preprocessing')
steps = [
    ('01', '시간 정렬', '전력수요와 7개 지역 ASOS를\ndatetime 기준으로 결합'),
    ('02', '결측 처리', '강수·일조·일사 = 0\n기온·풍속·습도 = 지역별 보간'),
    ('03', '주기 특성', 'hour / weekday / month를\nsin·cos 특성으로 변환'),
    ('04', '스케일링', 'Scaler는 학습 구간에만\n적합해 평가 정보 차단'),
]
for i, (num, head, body) in enumerate(steps):
    x = 0.85 + i * 3.0
    rect(slide, x, 1.75, 2.55, 3.45, PALE, PALE, True)
    text(slide, num, x + 0.22, 2.0, 0.55, 0.35, 18, CYAN, True)
    text(slide, head, x + 0.22, 2.55, 2.05, 0.35, 16, NAVY, True)
    text(slide, body, x + 0.22, 3.25, 2.1, 0.95, 13, DARK)
    if i < 3:
        text(slide, '→', x + 2.62, 3.1, 0.35, 0.4, 22, ORANGE, True, PP_ALIGN.CENTER)
text(slide, '결측을 무조건 보간하지 않고 변수의 물리적 의미에 맞춰 처리했습니다.', 0.9, 6.05, 8.8, 0.4, 15, NAVY, True)
footer(slide, 4)

# 5. Fair evaluation
slide = prs.slides.add_slide(blank)
set_bg(slide, NAVY)
title(slide, '공정한 시험: 모두 다음 1시간 예측', '04 / evaluation', WHITE)
text(slide, '70%', 0.95, 1.8, 2.1, 0.7, 38, CYAN, True)
text(slide, '학습 6,132시간', 1.0, 2.55, 2.4, 0.3, 13, WHITE, True)
text(slide, '15%', 4.35, 1.8, 2.1, 0.7, 38, ORANGE, True)
text(slide, '검증 1,314시간', 4.4, 2.55, 2.4, 0.3, 13, WHITE, True)
text(slide, '15%', 7.75, 1.8, 2.1, 0.7, 38, GREEN, True)
text(slide, '평가 1,314시간', 7.8, 2.55, 2.4, 0.3, 13, WHITE, True)
rect(slide, 0.95, 3.55, 10.2, 0.12, RGBColor(75, 91, 107))
rect(slide, 0.95, 3.55, 7.14, 0.12, CYAN)
rect(slide, 8.09, 3.55, 1.53, 0.12, ORANGE)
rect(slide, 9.62, 3.55, 1.53, 0.12, GREEN)
text(slide, 'LSTM / Transformer', 1.0, 4.25, 2.7, 0.35, 15, WHITE, True)
text(slide, '최근 24시간 실제 자료 → 다음 1시간', 1.0, 4.7, 4.2, 0.35, 13, RGBColor(203, 220, 228))
text(slide, 'ARIMA', 6.2, 4.25, 2.2, 0.35, 15, WHITE, True)
text(slide, '1시간 예측 → 실제값 반영 → 다음 1시간', 6.2, 4.7, 5.5, 0.35, 13, RGBColor(203, 220, 228))
text(slide, 'ARIMA만 수개월을 한 번에 예측하는 조건을 제거했습니다.', 1.0, 5.9, 8.4, 0.38, 17, CYAN, True)
footer(slide, 5)

# 6. Models
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, '비교한 세 가지 모델', '05 / models')
model_data = [
    ('ARIMA / SARIMAX', '고전 시계열 baseline', '추세와 24시간 계절성\n해석 가능한 기준선', BLUE),
    ('LSTM', '순환 신경망', '최근 24시간의\n장기 의존성 학습', ORANGE),
    ('Transformer', 'Self-attention', '시간적 관계와\n다변량 상호작용 학습', CYAN),
]
for i, (head, sub, body, accent) in enumerate(model_data):
    x = 0.85 + i * 4.1
    rect(slide, x, 1.8, 3.55, 3.8, PALE, PALE, True)
    rect(slide, x, 1.8, 3.55, 0.18, accent)
    text(slide, head, x + 0.28, 2.35, 3.0, 0.38, 19, NAVY, True)
    text(slide, sub, x + 0.28, 2.85, 3.0, 0.3, 11, accent, True)
    text(slide, body, x + 0.28, 3.55, 2.8, 0.75, 16, DARK)
text(slide, '공통 입력: 지역 날씨 + 시간 특성 + 과거 전력수요', 1.0, 6.2, 8.5, 0.35, 16, NAVY, True)
footer(slide, 6)

# 7. Results
slide = prs.slides.add_slide(blank)
set_bg(slide, WHITE)
title(slide, '성능 비교 결과', '06 / results')
rows = [['모델', 'MAE', 'RMSE', 'MAPE', '표준편차 비율', '95% 포함률']]
for _, r in metrics.iterrows():
    rows.append([r['model'], f"{r['MAE']:,.1f}", f"{r['RMSE']:,.1f}", f"{r['MAPE']:.2f}%", f"{r['std_ratio']:.3f}", f"{r['interval_coverage_95']:.2f}%"])
for i, row in enumerate(rows):
    y = 1.8 + i * 0.55
    fill = BLUE if i == 0 else (PALE if i % 2 else WHITE)
    rect(slide, 0.85, y, 11.8, 0.55, fill)
    for j, value in enumerate(row):
        widths = [2.1, 2.3, 2.3, 1.8, 3.0, 2.7]
        x = 0.85 + sum(widths[:j])
        text(slide, str(value), x, y + 0.09, widths[j], 0.3, 10, WHITE if i == 0 else DARK, i == 0 or j == 0, PP_ALIGN.CENTER if i == 0 else PP_ALIGN.RIGHT)
text(slide, '표준편차 비율이 0에 가깝지 않아 LSTM·Transformer의 상수 예측 붕괴는 관찰되지 않았습니다.', 0.95, 4.55, 11.2, 0.5, 16, NAVY, True)
add_image(slide, 'forecast_comparison.png', 0.95, 5.1, 5.75, 1.85)
text(slide, '그래프와 수치를 함께 읽어야 모델의 실제 성능을 판단할 수 있습니다.', 7.1, 5.45, 4.6, 0.75, 16, DARK, True)
footer(slide, 7)

# 8. Uncertainty
slide = prs.slides.add_slide(blank)
set_bg(slide)
title(slide, '예측구간과 불확실성', '07 / uncertainty')
add_image(slide, 'arima_interval.png', 0.75, 1.55, 3.95, 2.0)
add_image(slide, 'lstm_interval.png', 4.7, 1.55, 3.95, 2.0)
add_image(slide, 'transformer_interval.png', 8.65, 1.55, 3.95, 2.0)
text(slide, 'ARIMA', 2.1, 3.7, 1.3, 0.25, 11, BLUE, True, PP_ALIGN.CENTER)
text(slide, 'LSTM', 6.1, 3.7, 1.3, 0.25, 11, ORANGE, True, PP_ALIGN.CENTER)
text(slide, 'Transformer', 9.9, 3.7, 1.5, 0.25, 11, CYAN, True, PP_ALIGN.CENTER)
bullet_list(slide, ['95% 구간 폭은 검증 잔차 표준편차로 계산', '평가 정답을 구간 계산에 사용하지 않아 누출 방지', '포함률은 실제값이 구간 안에 들어간 비율', '현재 구간은 잔차 기반 근사이며 확률 보장은 아님'], 1.0, 4.45, 11.0, 15, DARK, 0.48)
footer(slide, 8)

# 9. Interpretation
slide = prs.slides.add_slide(blank)
set_bg(slide, PALE)
title(slide, '결과 해석', '08 / interpretation')
rect(slide, 0.85, 1.7, 5.65, 4.5, WHITE, WHITE, True)
text(slide, '이번 실험에서 확인한 것', 1.2, 2.05, 4.3, 0.35, 17, NAVY, True)
bullet_list(slide, ['7개 지역 날씨를 전국 수요 예측에 함께 사용', '세 모델의 시험 조건을 one-step 방식으로 통일', '입력 스케일러는 학습 데이터에만 적합', 'LSTM·Transformer의 예측 변동성은 실제와 유사'], 1.2, 2.75, 4.75, 14, DARK, 0.63)
rect(slide, 6.9, 1.7, 5.55, 4.5, NAVY, NAVY, True)
text(slide, '남은 한계와 다음 단계', 7.3, 2.05, 4.5, 0.35, 17, WHITE, True)
bullet_list(slide, ['ARIMA의 수렴과 차수 최적화', 'quantile loss 기반 예측구간', 'conformal prediction으로 coverage 보정', '더 긴 기간과 다른 연도 데이터 검증'], 7.3, 2.75, 4.55, 14, RGBColor(220, 235, 241), 0.63)
footer(slide, 9)

# 10. Closing
slide = prs.slides.add_slide(blank)
set_bg(slide, NAVY)
text(slide, 'THANK YOU', 0.85, 1.0, 4.2, 0.4, 13, CYAN, True)
text(slide, '전력수요 예측은\n정확도만의 문제가 아니다.', 0.85, 2.0, 7.4, 1.25, 32, WHITE, True)
text(slide, '공정한 시험 조건, 지역 기상 정보,\n그리고 예측의 불확실성까지 함께 보아야 한다.', 0.9, 4.05, 7.5, 0.8, 18, RGBColor(213, 229, 236))
text(slide, '전국 전력수요 시계열 예측  /  주희', 0.9, 6.25, 6.0, 0.3, 12, CYAN, True)
footer(slide, 10)

prs.save(OUTPUT)
print(OUTPUT)
